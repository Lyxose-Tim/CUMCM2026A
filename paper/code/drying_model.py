# -*- coding: utf-8 -*-
"""
drying_model.py —— 药材热风烘干的一维径向传热传质模型（2026 年 A 题，问题一至四）

用法：python drying_model.py [--data 附件目录] [--out 输出目录]
输出：result1.xlsx～result4.xlsx，以及论文表 1～7 的数值（table*.csv）。
依赖：numpy、scipy、openpyxl。

模型：干基含水率 C 与温度 T 满足径向散度形式的扩散方程与导热方程，
      表面为对流传质、对流换热的第三类边界，中心对称；
      空间用节点型有限体积法（含水率界面系数取积分平均，导热系数取调和平均），
      时间用自适应 BDF（scipy.integrate.solve_ivp），烘干终止判据为全域最大含水率降到 0.15。
问题四在参考坐标 x = r/R(t) 上求解，R(t) 取自附件 2。
"""
from __future__ import annotations

import argparse
import os
from pathlib import Path

import numpy as np
import openpyxl
from numpy.polynomial.legendre import leggauss
from openpyxl.cell import WriteOnlyCell
from scipy.integrate import solve_ivp
from scipy.sparse import lil_matrix

# ============================== 1 题给参数与数值设置 ==============================
R0 = 0.02                    # 初始半径 / m
T0_C = 28.0                  # 初始温度 / ℃
C0 = 2.55                    # 初始干基含水率 / (kg/kg)
H = 25.0                     # 对流换热系数 / (W/(m²·K))，附录 2，各问沿用
HM = 8.0e-7                  # 对流传质系数 / (m/s)，附录 2，各问沿用
THRESH = 0.15                # 烘干要求：药材各处含水率低于 0.15
KELVIN = 273.15
T_SWITCH = 14400.0           # 附件 1 数据终点（4 h），其后取末小时均值
WINDOW = (10800, 14400)      # 末小时窗口（含两端，61 个样本）
N_FINAL = {"q1": 1600, "q23": 800, "q4": 800}      # 各问计算网格区间数
NPTS = 8                     # 积分平均界面系数的 Gauss 点数
RTOL, ATOL_C, ATOL_T = 1e-8, 1e-11, 1e-8           # BDF 容差
MAXSTEP_DATA, MAXSTEP_AFTER = 60.0, 600.0          # 最大步长：4 h 内 / 4 h 后
POST_S = 600.0               # 阈值穿越后的续算检查时长 / s
T_CAP_H = {"q23": 90.0, "q4": 200.0}               # 积分上限 / h
COLS21 = [round(0.1 * j, 4) for j in range(21)]    # 0.0, 0.1, …, 2.0 cm
COLS20 = [round(0.1 * j, 4) for j in range(20)]    # 0.0, 0.1, …, 1.9 cm（问题四另加表面列）
A1 = "时间\\到药材中心的距离"


# ============================== 2 经验物性（附录 2/3/4） ==============================
def _exp_neg_over(beta, C):
    """exp(-beta/C)：C≤0 处取 0；指数小于 -700 时以 0 近似（避免下溢）。"""
    C = np.asarray(C, dtype=np.float64)
    out = np.zeros_like(C)
    pos = C > 0.0
    if np.any(pos):
        arg = -beta / C[pos]
        keep = arg >= -700.0
        vals = np.zeros_like(arg)
        vals[keep] = np.exp(arg[keep])
        out[pos] = vals
    return out


def _exp_neg_over_T(gamma, T):
    """exp(-gamma/T)，T 为热力学温度 / K。"""
    T = np.asarray(T, dtype=np.float64)
    arg = -gamma / T
    out = np.zeros_like(arg)
    keep = arg >= -700.0
    out[keep] = np.exp(arg[keep])
    return out


class PropsQ1:
    """附录 2：常热物性，扩散系数只依赖含水率。"""
    def rho(self, C):
        C = np.asarray(C, dtype=np.float64)
        return np.full_like(C, 820.0)

    def cp(self, C):
        C = np.asarray(C, dtype=np.float64)
        return np.full_like(C, 2600.0)

    def k(self, C):
        C = np.asarray(C, dtype=np.float64)
        return np.full_like(C, 0.36)

    def b(self, C):
        return self.rho(C) * self.cp(C)

    def D(self, C, T=None):
        return 7.0e-9 * _exp_neg_over(0.89, C)


class PropsQ23:
    """附录 3：问题二、三的变物性关系。"""
    def rho(self, C):
        C = np.asarray(C, dtype=np.float64)
        return 650.0 + 128.0 * C

    def cp(self, C):
        C = np.asarray(C, dtype=np.float64)
        return 1450.0 + 2736.0 * C / (C + 1.0)

    def k(self, C):
        C = np.asarray(C, dtype=np.float64)
        return 0.21 + 0.38 * C / (C + 1.0)

    def b(self, C):
        return self.rho(C) * self.cp(C)

    def D(self, C, T):
        return 2.4e-3 * _exp_neg_over(0.45, C) * _exp_neg_over_T(3850.0, T)


class PropsQ4:
    """附录 4：问题四的变物性关系。"""
    def rho(self, C):
        C = np.asarray(C, dtype=np.float64)
        return 760.0 + 90.0 * C

    def cp(self, C):
        C = np.asarray(C, dtype=np.float64)
        return 1850.0 + 2150.0 * C / (C + 1.0)

    def k(self, C):
        C = np.asarray(C, dtype=np.float64)
        return 0.12 + 0.20 * C / (C + 1.0)

    def b(self, C):
        return self.rho(C) * self.cp(C)

    def D(self, C, T):
        return 4.2e-4 * _exp_neg_over(0.30, C) * _exp_neg_over_T(3850.0, T)


PROPS = {"q1": PropsQ1, "q23": PropsQ23, "q4": PropsQ4}

_U, _W = leggauss(NPTS)
_XI, _WT = 0.5 * (_U + 1.0), 0.5 * _W      # [0,1] 上的 Gauss 节点与权重


def face_integral(Ci, Cj, Tb, Dfun):
    """积分平均界面系数 ∫_0^1 D((1-ξ)Ci+ξCj, Tb) dξ（8 点 Gauss）。"""
    acc = np.zeros(np.broadcast(Ci, Cj, Tb).shape, dtype=np.float64)
    for xk, wk in zip(_XI, _WT):
        acc = acc + wk * Dfun((1.0 - xk) * Ci + xk * Cj, Tb)
    return acc


def face_harmonic(a, b):
    """调和平均界面系数 2ab/(a+b)。"""
    s = a + b
    out = np.zeros_like(s)
    nz = s > 0.0
    out[nz] = 2.0 * a[nz] * b[nz] / s[nz]
    return out


# ============================== 3 环境与半径输入 ==============================
def read_columns(path):
    """读取附件（首行为表头）的数值列。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [r for i, r in enumerate(ws.iter_rows(values_only=True)) if i > 0 and r[0] is not None]
    wb.close()
    return [np.array([float(r[j]) for r in rows]) for j in range(len(rows[0]))]


class Env:
    """烘房空气温度 T_air(t)/K 与环境含水浓度 C_env(t)：4 h 内线性插值，其后取常值。"""
    def __init__(self, t_nodes, T_nodes_K, C_nodes, T_const_K, C_const):
        self.t_nodes, self.T_nodes_K, self.C_nodes = t_nodes, T_nodes_K, C_nodes
        self.T_const_K, self.C_const = T_const_K, C_const

    def T_air_K(self, t):
        t = np.asarray(t, dtype=np.float64)
        out = np.where(t <= T_SWITCH, np.interp(t, self.t_nodes, self.T_nodes_K), self.T_const_K)
        return out if out.shape else float(out)

    def C_env(self, t):
        t = np.asarray(t, dtype=np.float64)
        out = np.where(t <= T_SWITCH, np.interp(t, self.t_nodes, self.C_nodes), self.C_const)
        return out if out.shape else float(out)


def load_env_data(data_dir):
    """附件 1：返回 (t/s, T_air/℃, C_env)。"""
    t, T_C, C = read_columns(Path(data_dir) / "附件1.xlsx")
    return t, T_C, C


def make_env(t, T_C, C):
    """基准环境：原始节点线性插值，4 h 后取末小时 61 个样本的均值。"""
    win = (t >= WINDOW[0]) & (t <= WINDOW[1])
    return Env(t, T_C + KELVIN, C.copy(), float(np.mean(T_C[win])) + KELVIN, float(np.mean(C[win])))


class Radius:
    """附件 2 的半径 R(t)/m：线性插值，72 h 后保持末值。"""
    def __init__(self, t_nodes, R_nodes_m):
        self.t_nodes, self.R_nodes_m = t_nodes, R_nodes_m

    def R(self, t):
        t = np.asarray(t, dtype=np.float64)
        out = np.interp(t, self.t_nodes, self.R_nodes_m)
        return out if out.shape else float(out)


def load_radius(data_dir):
    t, R_cm = read_columns(Path(data_dir) / "附件2.xlsx")
    return Radius(t, R_cm / 100.0)


# ============================== 4 节点型有限体积离散 ==============================
class Grid:
    """节点 i·d（i=0..N），界面 (i+1/2)d，控制体权重 V_i = ∫ r dr（约去 2πL）。"""
    def __init__(self, N, length):
        self.N, self.d = N, length / N
        self.x = np.arange(N + 1) * self.d
        self.A = (np.arange(N) + 0.5) * self.d
        V = np.empty(N + 1)
        V[0] = self.d ** 2 / 8.0
        V[1:N] = self.x[1:N] * self.d
        V[N] = length * self.d / 2.0 - self.d ** 2 / 8.0
        self.V = V                                   # ΣV_i = length²/2


class Model:
    """状态 y=[C_0..C_N, T_0..T_N]。radius=None 为固定域（问题一至三），否则为参考坐标（问题四）。"""
    def __init__(self, N, props, env, *, radius=None, R0=R0, h=H, hm=HM):
        self.N, self.props, self.env, self.radius = N, props, env, radius
        self.ref = radius is not None
        self.R0, self.h, self.hm = float(R0), float(h), float(hm)
        self.grid = Grid(N, 1.0 if self.ref else self.R0)

    def surf(self, t):
        """表面交换系数：固定域 (R0·h_m, R0·h)，参考坐标 (h_m/R, h/R)。"""
        if self.ref:
            R = float(self.radius.R(t))
            return self.hm / R, self.h / R
        return self.R0 * self.hm, self.R0 * self.h

    def scale(self, t):
        """扩散项几何因子：参考坐标为 1/R(t)²。"""
        if self.ref:
            R = float(self.radius.R(t))
            return 1.0 / (R * R)
        return 1.0

    def rhs(self, t, y):
        N, g = self.N, self.grid
        n = N + 1
        C, T = y[:n], y[n:2 * n]
        Cenv, Tair = float(self.env.C_env(t)), float(self.env.T_air_K(t))
        sC, sT = self.surf(t)
        Dface = face_integral(C[:-1], C[1:], 0.5 * (T[:-1] + T[1:]), self.props.D)
        wC = g.A * Dface / g.d * self.scale(t)                       # 含水率传输率
        kn = self.props.k(C)
        wT = g.A * face_harmonic(kn[:-1], kn[1:]) / g.d * self.scale(t)   # 导热传输率

        dC = np.empty(n)
        dC[1:N] = (wC[:N - 1] * (C[:N - 1] - C[1:N]) + wC[1:N] * (C[2:N + 1] - C[1:N])) / g.V[1:N]
        dC[0] = wC[0] * (C[1] - C[0]) / g.V[0]
        dC[N] = (wC[N - 1] * (C[N - 1] - C[N]) - sC * (C[N] - Cenv)) / g.V[N]

        b = self.props.b(C)
        dT = np.empty(n)
        dT[1:N] = (wT[:N - 1] * (T[:N - 1] - T[1:N]) + wT[1:N] * (T[2:N + 1] - T[1:N])) / (b[1:N] * g.V[1:N])
        dT[0] = wT[0] * (T[1] - T[0]) / (b[0] * g.V[0])
        dT[N] = (wT[N - 1] * (T[N - 1] - T[N]) - sT * (T[N] - Tair)) / (b[N] * g.V[N])
        return np.concatenate([dC, dT])

    def initial_state(self):
        n = self.N + 1
        return np.concatenate([np.full(n, C0), np.full(n, T0_C + KELVIN)])

    def cbar(self, C):
        """干物质质量加权平均含水率：固定域 (2/R0²)ΣV_iC_i，参考坐标 2ΣV_ic_i。"""
        L = 1.0 if self.ref else self.R0
        return float(np.dot(2.0 / L ** 2 * self.grid.V, C))

    def drying_rate(self, t, C):
        """平均含水率的下降速率 -dC̄/dt = (2h_m/R)(C_s − C_env)。"""
        R = float(self.radius.R(t)) if self.ref else self.R0
        return 2.0 * self.hm / R * (float(C[-1]) - float(self.env.C_env(t)))


def sparsity(N):
    """Jacobian 稀疏结构：2×2 块三对角（含 C、T 交叉耦合）。"""
    n = N + 1
    S = lil_matrix((2 * n, 2 * n), dtype=np.int8)
    for i in range(n):
        for j in (i - 1, i, i + 1):
            if 0 <= j < n:
                S[i, j] = S[i, n + j] = S[n + i, n + j] = S[n + i, j] = 1
    return S.tocsr()


# ============================== 5 自适应 BDF 时间积分 ==============================
class Trajectory:
    """分段 BDF 解（含稠密输出）。eval 只在已积分区间内求值。"""
    def __init__(self, segments, t_end, t_cross=None, y_cross=None):
        self.segments, self.t_end, self.t_cross, self.y_cross = segments, t_end, t_cross, y_cross

    def eval(self, ts):
        ts = np.atleast_1d(np.asarray(ts, dtype=np.float64))
        out = np.empty((ts.size, self.segments[0].y.shape[0]))
        for k, ti in enumerate(ts):
            for seg in self.segments:
                if seg.t[0] - 1e-6 <= ti <= seg.t[-1] + 1e-6:
                    out[k] = seg.sol(ti)
                    break
            else:
                raise ValueError(f"t={ti} 超出已积分区间")
        return out


def integrate(model, y0, t0, t_end, *, breakpoints=(T_SWITCH,), threshold=None, rtol=RTOL):
    """从 t0 积分到 t_end；在 4 h 断点重启；给 threshold 时以 max C 下穿阈值为终止事件。"""
    n = model.N + 1
    atol = np.concatenate([np.full(n, ATOL_C), np.full(n, ATOL_T)])
    S = sparsity(model.N)
    edges = [t0] + sorted(b for b in breakpoints if t0 < b < t_end) + [t_end]
    event = None
    if threshold is not None:
        def event(t, y):
            return float(np.max(y[:n])) - threshold
        event.direction, event.terminal = -1.0, True
    segs, y = [], np.asarray(y0, dtype=np.float64).copy()
    for a, b in zip(edges[:-1], edges[1:]):
        max_step = MAXSTEP_DATA if b <= T_SWITCH + 1e-9 else MAXSTEP_AFTER
        sol = solve_ivp(model.rhs, (a, b), y, method="BDF", rtol=rtol, atol=atol,
                        jac_sparsity=S, dense_output=True, max_step=max_step, events=event)
        if not sol.success:
            raise RuntimeError(f"BDF 在 [{a}, {b}] 失败：{sol.message}")
        segs.append(sol)
        y = sol.y[:, -1].copy()
        if event is not None and len(sol.t_events[0]) > 0:
            return Trajectory(segs, sol.t[-1], float(sol.t_events[0][0]), sol.y_events[0][0].copy())
    return Trajectory(segs, t_end)


def solve_to_dry(model, t_cap_h, *, rtol=RTOL):
    """积分至全域最大含水率降到 0.15（阈值穿越时刻 t*），再续算 600 s，拼成同一条轨迹。
    返回 (轨迹, t*, 首个严格低于阈值的 60 s 采样时刻, 续算段最大含水率)。"""
    n = model.N + 1
    res = integrate(model, model.initial_state(), 0.0, t_cap_h * 3600.0,
                    breakpoints=(T_SWITCH,), threshold=THRESH, rtol=rtol)
    if res.t_cross is None:
        raise RuntimeError(f"{t_cap_h} h 内未达到烘干要求")
    cont = integrate(model, res.y_cross, res.t_cross, res.t_cross + POST_S,
                     breakpoints=(), threshold=None, rtol=rtol)

    def cmax_cont(t):
        return float(np.max(cont.eval([t])[0][:n]))

    t_sample, t = None, int(np.ceil(res.t_cross / 60.0)) * 60.0
    while t <= res.t_cross + POST_S + 1e-9:
        if cmax_cont(t) < THRESH:
            t_sample = float(t)
            break
        t += 60.0
    post_max = max(cmax_cont(t) for t in np.linspace(res.t_cross, res.t_cross + POST_S, 61))
    full = Trajectory(res.segments + cont.segments, cont.t_end, res.t_cross, res.y_cross)
    return full, res.t_cross, t_sample, post_max


# ============================== 6 结果文件与表格 ==============================
def _cell(ws, v):
    if v is None:
        return None
    c = WriteOnlyCell(ws, value=float(np.round(v, 4)))
    c.number_format = "0.0000"
    return c


def write_result12(path, ts, C, T_C):
    """result1/2：工作表「温度」「水分浓度」。"""
    wb = openpyxl.Workbook(write_only=True)
    for name, grid in (("温度", T_C), ("水分浓度", C)):
        ws = wb.create_sheet(name)
        ws.append([A1] + [float(c) for c in COLS21])
        for i, t in enumerate(ts):
            ws.append([int(round(t))] + [_cell(ws, grid[i][j]) for j in range(len(COLS21))])
    wb.save(path)


def write_result34(path, ts, rows, cols, surface=None):
    """result3/4：单工作表；问题四另加「药材表面」列，域外单元格留空。"""
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Sheet1")
    ws.append([A1] + [float(c) for c in cols] + (["药材表面"] if surface is not None else []))
    for i, t in enumerate(ts):
        row = [int(round(t))] + [_cell(ws, rows[i][j]) for j in range(len(cols))]
        if surface is not None:
            row.append(_cell(ws, surface[i]))
        ws.append(row)
    wb.save(path)


def write_csv(path, header, rows):
    with open(path, "w", encoding="utf-8") as f:
        f.write(",".join(header) + "\n")
        for r in rows:
            f.write(",".join("" if v is None else (f"{v:.4f}" if isinstance(v, float) else str(v))
                             for v in r) + "\n")


def node_index(model, r_cm):
    return int(round(r_cm / 100.0 / model.grid.d))


def q4_row(c, model, R_t, cols):
    """问题四：固定物理位置 r_j 映射到 x_j=r_j/R(t) 后线性插值；位于表面之外的位置留空。"""
    row = []
    for r_cm in cols:
        r_m = r_cm / 100.0
        if not r_m <= R_t * (1.0 + 1e-12):
            row.append(None)
            continue
        xj = r_m / R_t
        row.append(float(c[-1]) if xj >= 1.0 - 1e-12 else float(np.interp(xj, model.grid.x, c)))
    return row


def table_hours(t_cross, step_h=6.0):
    """每隔 6 h 一行，末行为烘干完成时刻 t*。"""
    times = list(np.arange(step_h, t_cross / 3600.0 + 1e-9, step_h) * 3600.0)
    if not times or abs(times[-1] - t_cross) > 1.0:
        times.append(t_cross)
    else:
        times[-1] = t_cross
    return times


def produce_q1(env, out):
    """问题一：0～1800 s，每 1 s、每 0.1 cm（附录 2，N=1600）。"""
    m = Model(N_FINAL["q1"], PropsQ1(), env)
    n = m.N + 1
    res = integrate(m, m.initial_state(), 0.0, 1800.0, breakpoints=())
    idx = np.array([node_index(m, r) for r in COLS21])
    ts = np.arange(1, 1801)
    Y = res.eval(ts)
    write_result12(out / "result1.xlsx", ts, Y[:, idx], Y[:, idx + n] - KELVIN)
    tt = [100, 300, 600, 900, 1200, 1500, 1800]
    i5 = np.array([node_index(m, r) for r in (0.0, 0.5, 1.0, 1.5, 2.0)])
    Yt = res.eval(tt)
    hdr = ["t", "r0", "r0.5", "r1.0", "r1.5", "r2.0"]
    write_csv(out / "table1_temp.csv", hdr, [[t] + list(Yt[k, i5 + n] - KELVIN) for k, t in enumerate(tt)])
    write_csv(out / "table2_moist.csv", hdr, [[t] + list(Yt[k, i5]) for k, t in enumerate(tt)])


def produce_q23(env, out):
    """问题二、三：同一条轨迹给出 result2（每 1 s）与 result3（每 60 s）及表 3～5（附录 3，N=800）。"""
    m = Model(N_FINAL["q23"], PropsQ23(), env)
    n = m.N + 1
    full, t_cross, t_sample, post_max = solve_to_dry(m, T_CAP_H["q23"])
    t_end_1s = int(np.ceil(t_cross))
    while float(np.max(full.eval([float(t_end_1s)])[0][:n])) >= THRESH:
        t_end_1s += 1
    idx = np.array([node_index(m, r) for r in COLS21])
    ts2 = np.arange(1, min(t_end_1s, int(full.t_end)) + 1)
    C2, T2 = [], []
    for i in range(0, len(ts2), 8000):
        Y = full.eval(ts2[i:i + 8000])
        C2.append(Y[:, idx].copy())
        T2.append(Y[:, idx + n].copy())
    write_result12(out / "result2.xlsx", ts2, np.vstack(C2), np.vstack(T2) - KELVIN)
    ts3 = np.arange(60, int(t_sample) + 1, 60)
    write_result34(out / "result3.xlsx", ts3, full.eval(ts3)[:, idx], COLS21)
    i5 = np.array([node_index(m, r) for r in (0.0, 0.5, 1.0, 1.5, 2.0)])
    t34 = np.arange(0.5, 3.01, 0.5) * 3600.0
    Y34 = full.eval(t34)
    hdr = ["t", "r0", "r0.5", "r1.0", "r1.5", "r2.0"]
    write_csv(out / "table3_temp.csv", hdr, [[f"{t / 3600:g}"] + list(Y34[k, i5 + n] - KELVIN) for k, t in enumerate(t34)])
    write_csv(out / "table4_moist.csv", hdr, [[f"{t / 3600:g}"] + list(Y34[k, i5]) for k, t in enumerate(t34)])
    write_csv(out / "table5_moist.csv", ["t_h", "r0", "r0.5", "r1.0", "r1.5", "r2.0"],
              [[t / 3600] + [float(v) for v in full.eval([t])[0][i5]] for t in table_hours(t_cross)])
    return {"t_star_s": t_cross, "t_end_1s": t_end_1s, "t_sample_s": t_sample, "post_max": post_max}


def produce_q4(env, radius, out):
    """问题四：参考坐标上求解（附录 4，N=800），每 60 s 输出固定位置与药材表面的含水率。"""
    m = Model(N_FINAL["q4"], PropsQ4(), env, radius=radius)
    n = m.N + 1
    full, t_cross, t_sample, post_max = solve_to_dry(m, T_CAP_H["q4"])
    ts4 = np.arange(60, int(t_sample) + 1, 60)
    rows, surf = [], []
    for t in ts4:
        c = full.eval([float(t)])[0][:n]
        rows.append(q4_row(c, m, float(radius.R(float(t))), COLS20))
        surf.append(float(c[-1]))
    write_result34(out / "result4.xlsx", ts4, rows, COLS20, surface=surf)
    t6, rad = [], []
    for t in table_hours(t_cross):
        c = full.eval([t])[0][:n]
        R_t = float(radius.R(t))
        r0 = q4_row(c, m, R_t, [0.0, 0.5, 1.0])
        t6.append([t / 3600, r0[0], r0[1], r0[2], float(c[-1])])
        rad.append([t / 3600, R_t * 100, int(t > radius.t_nodes[-1])])
    write_csv(out / "table6_moist.csv", ["t_h", "r0", "r0.5", "r1.0", "surface"], t6)
    write_csv(out / "table6_radius.csv", ["t_h", "R_cm", "extrapolated"], rad)
    return {"t_star_s": t_cross, "t_sample_s": t_sample, "post_max": post_max}


def main(argv=None):
    here = Path(__file__).resolve().parent
    ap = argparse.ArgumentParser(description="药材烘干模型：生成 result1～4 与论文表格")
    ap.add_argument("--data", default=str(here.parent.parent / "附件"), help="附件1/附件2 所在目录")
    ap.add_argument("--out", default=str(here / "output"), help="输出目录")
    a = ap.parse_args(argv)
    out = Path(a.out)
    os.makedirs(out, exist_ok=True)
    t, T_C, C = load_env_data(a.data)
    env, radius = make_env(t, T_C, C), load_radius(a.data)
    produce_q1(env, out)
    r23 = produce_q23(env, out)
    r4 = produce_q4(env, radius, out)
    print(f"问题三 烘干时长 t* = {r23['t_star_s'] / 3600:.4f} h（{r23['t_star_s']:.6f} s）")
    print(f"问题四 烘干时长 t* = {r4['t_star_s'] / 3600:.4f} h（{r4['t_star_s']:.6f} s）")


if __name__ == "__main__":
    main()
