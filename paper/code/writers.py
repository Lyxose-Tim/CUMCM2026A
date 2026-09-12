"""writers.py —— result1–4 写盘与写后读回校验（§9.5、V-9）。

- openpyxl write_only 逐行写入，数值型 + `0.0000` 格式，仅写盘时 round(·,4)。
- result1/2：工作表 温度/水分浓度；result3/4：Sheet1。A1=`时间\\到药材中心的距离`。
- result4：列 0..1.9 + 「药材表面」；域外单元格留空（None→空）。
- 数据行 ≤1,048,575（含表头 ≤1,048,576），超限抛错，不截断、不改模板。
- verify_workbook：读回校验表名/A1/表头/A 列连续无重复/行数/掩码。
"""
from __future__ import annotations

from pathlib import Path

import numpy as np
import openpyxl
from openpyxl.cell import WriteOnlyCell

from .postprocess import round4

EXCEL_MAX_ROWS = 1048576
NUMFMT = "0.0000"


def _num_cell(ws, value):
    """数值型单元格（round4 + 0.0000 格式）；None → 空单元格。"""
    if value is None:
        return None
    c = WriteOnlyCell(ws, value=round4(value))
    c.number_format = NUMFMT
    return c


def _check_row_budget(n_data_rows: int, path):
    if n_data_rows + 1 > EXCEL_MAX_ROWS:
        raise ValueError(f"{path}: 数据行 {n_data_rows} 超出上限 {EXCEL_MAX_ROWS - 1}")


def write_result12(path, t_array, C_grid, T_grid, cols_cm, a1_text):
    """result1/2：两工作表（温度、水分浓度）。C_grid/T_grid 形状 (len(t), len(cols))。"""
    path = Path(path)
    t_array = np.asarray(t_array)
    _check_row_budget(len(t_array), path)
    wb = openpyxl.Workbook(write_only=True)
    for sheet_name, grid, is_temp in (("温度", T_grid, True), ("水分浓度", C_grid, False)):
        ws = wb.create_sheet(sheet_name)
        ws.append([a1_text] + [float(c) for c in cols_cm])
        for i, t in enumerate(t_array):
            row = [int(round(t))] + [_num_cell(ws, grid[i][j]) for j in range(len(cols_cm))]
            ws.append(row)
    wb.save(path)


def write_result34(path, t_array, value_grid, cols_cm, a1_text, *,
                   sheet_name="Sheet1", surface_header=None, surface_values=None):
    """result3/4：单工作表（水分浓度）。value_grid 可含 None（域外留空）。

    result4 传 surface_header='药材表面' 与 surface_values（长度 len(t)）。
    """
    path = Path(path)
    t_array = np.asarray(t_array)
    _check_row_budget(len(t_array), path)
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet(sheet_name)
    header = [a1_text] + [float(c) for c in cols_cm]
    if surface_header is not None:
        header.append(surface_header)
    ws.append(header)
    for i, t in enumerate(t_array):
        row = [int(round(t))] + [_num_cell(ws, value_grid[i][j]) for j in range(len(cols_cm))]
        if surface_header is not None:
            row.append(_num_cell(ws, surface_values[i]))
        ws.append(row)
    wb.save(path)


def verify_workbook(path, *, expected_sheets, a1_text, expected_cols, n_data_rows,
                    t_start, t_step, surface_header=None, mask=None,
                    full_format_check=False, require_surface_nonempty=False):
    """写后读回校验（V-9）。mask[i][j]=True 表域内（应非空）。

    full_format_check=True：检查**所有数据行**的数字格式（生产用，非抽查）。
    require_surface_nonempty=True：result4 表面列（末列）不得为空。返回检查结果 dict。
    """
    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True)
    issues = []
    if list(wb.sheetnames) != list(expected_sheets):
        issues.append(f"工作表名 {wb.sheetnames} != {expected_sheets}")
    exp_header = [a1_text] + [float(c) for c in expected_cols] + \
                 ([surface_header] if surface_header else [])
    width = len(exp_header)
    for sn in wb.sheetnames:
        ws = wb[sn]
        # openpyxl write_only 会丢弃行尾 None（域外单元格），读回时补齐到表头宽度
        rows = [tuple(list(r) + [None] * (width - len(r))) if len(r) < width else r
                for r in ws.iter_rows(values_only=True)]
        header = list(rows[0])
        # 逐项表头（A1 + 数值列顺序 + 表面列）
        if header[0] != a1_text:
            issues.append(f"{sn}: A1={header[0]!r} != {a1_text!r}")
        if len(header) != len(exp_header):
            issues.append(f"{sn}: 表头列数 {len(header)} != {len(exp_header)}")
        else:
            for j, (g, e) in enumerate(zip(header[1:], exp_header[1:]), start=1):
                if isinstance(e, float):
                    if not (isinstance(g, (int, float)) and abs(float(g) - e) < 1e-9):
                        issues.append(f"{sn}: 表头列{j}={g!r} != {e}")
                elif g != e:
                    issues.append(f"{sn}: 表头列{j}={g!r} != {e!r}")
        # A 列连续整数、无重复、步距
        a_col = [r[0] for r in rows[1:]]
        if len(a_col) != n_data_rows:
            issues.append(f"{sn}: 数据行 {len(a_col)} != {n_data_rows}")
        expected_a = list(range(int(t_start), int(t_start) + n_data_rows * int(t_step), int(t_step)))
        if a_col != expected_a:
            issues.append(f"{sn}: A 列非预期连续整数步距 {t_step}（首 {a_col[:3]} 末 {a_col[-3:]}）")
        if len(set(a_col)) != len(a_col):
            issues.append(f"{sn}: A 列有重复")
        if not all(isinstance(v, int) and not isinstance(v, bool) for v in a_col):
            issues.append(f"{sn}: A 列须为整数秒")
        # 数据单元格：数值型（或 None 表域外），非字符串
        for i, r in enumerate(rows[1:]):
            for j, v in enumerate(r[1:], start=1):
                if v is not None and not (isinstance(v, (int, float)) and not isinstance(v, bool)):
                    issues.append(f"{sn}: 行{i} 列{j} 非数值型 {v!r}")
                    break
        # 数字格式 0.0000：full_format_check=True 单次流式检查所有数据行；否则抽查首/中/末
        if n_data_rows > 0:
            bad_fmt = False
            if full_format_check:
                rr = 1
                for cellrow in ws.iter_rows(min_row=2, max_row=n_data_rows + 1):
                    rr += 1
                    for col, cell in enumerate(cellrow[1:len(exp_header)], start=2):
                        if cell.value is not None and cell.number_format != NUMFMT:
                            issues.append(f"{sn}: 行{rr} 列{col} 数字格式 {cell.number_format!r} != {NUMFMT!r}")
                            bad_fmt = True
                            break
                    if bad_fmt:
                        break
            else:
                for rr in sorted({2, 2 + n_data_rows // 2, n_data_rows + 1}):
                    for col in range(2, len(exp_header) + 1):
                        cell = ws.cell(row=rr, column=col)
                        if cell.value is not None and cell.number_format != NUMFMT:
                            issues.append(f"{sn}: 行{rr} 列{col} 数字格式 {cell.number_format!r} != {NUMFMT!r}")
                            bad_fmt = True
                            break
                    if bad_fmt:
                        break
        # 表面列不得为空（result4）
        if require_surface_nonempty and surface_header is not None:
            scol = len(exp_header)     # 表面列为最后一列
            for i, r in enumerate(rows[1:]):
                if r[scol - 1] is None:
                    issues.append(f"{sn}: 行{i} 表面列为空（药材表面列不得留空）")
                    break
        # 域内空值：无掩码文件（result1/2/3）任何数据单元格不得为空
        if mask is None:
            for i, r in enumerate(rows[1:]):
                for j, v in enumerate(r[1:1 + len(expected_cols)]):
                    if v is None:
                        issues.append(f"{sn}: 行{i} 列{j} 域内空值（无掩码文件不得留空）")
                        break
        # 掩码一致性（result4）：域内非空、域外空 ⇔ inside()
        if mask is not None:
            for i, r in enumerate(rows[1:]):
                data = r[1:1 + len(expected_cols)]
                for j, v in enumerate(data):
                    if bool(mask[i][j]) == (v is None):
                        kind = "域内空值" if mask[i][j] else "域外非空"
                        issues.append(f"{sn}: 行{i} 列{j} 掩码不一致（{kind}）")
                        break
    wb.close()
    return {"ok": len(issues) == 0, "issues": issues}


def cross_file_check(result2_path, result3_path, *, decimals=4):
    """V-8 跨文件一致性：result2（1 s）与 result3（60 s）在共同 60 s 倍数时刻×位置
    舍入后相等（同源解）。返回不一致列表。"""
    def load_sheet(path, sheet):
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb[sheet] if sheet in wb.sheetnames else wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        header = rows[0]
        data = {int(r[0]): r[1:] for r in rows[1:]}
        return header, data

    h2, d2 = load_sheet(result2_path, "水分浓度")
    h3, d3 = load_sheet(result3_path, "Sheet1")
    issues = []
    if list(h2)[:len(h3)] != list(h3):
        # 表头前若干列应一致（result3 无温度表；比较水分列头）
        pass
    common = sorted(set(d2) & set(d3))
    common = [t for t in common if t % 60 == 0]
    # 覆盖缺口：result3 的 60 s 时刻若 ≤ result2 末时刻却不在 result2 → 应有共同时间未覆盖
    t2_max = max(d2)
    missing = sorted(t for t in d3 if t % 60 == 0 and t <= t2_max and t not in d2)
    if missing:
        issues.append(f"跨文件缺口：result3 的 60 s 时刻 {missing[:5]}… 在 result2 范围内却未被覆盖")
    # 零共同时刻不能判通过（预期应有共同时间）
    if len(common) == 0:
        issues.append("跨文件零共同时刻：result2 与 result3 无 60 s 公共采样时刻（应有共同时间，不判通过）")
    ncol = min(len(next(iter(d2.values()))), len(next(iter(d3.values()))))
    mism = 0
    for t in common:
        for j in range(ncol):
            v2, v3 = d2[t][j], d3[t][j]
            if v2 is None or v3 is None:
                continue
            if round(float(v2), decimals) != round(float(v3), decimals):
                mism += 1
                if len(issues) < 10:
                    issues.append(f"t={t}s 列{j}: result2={v2} != result3={v3}")
    ok = (mism == 0) and (not missing) and (len(common) > 0)
    return {"ok": ok, "n_common_times": len(common),
            "n_mismatch": mism, "n_missing_coverage": len(missing), "issues": issues}
