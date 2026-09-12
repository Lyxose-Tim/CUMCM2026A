"""writers.py —— result1–4 写盘与写后读回校验（§9.5、V-9）。

- openpyxl write_only 逐行写入，数值型 + `0.0000` 格式，仅写盘时 round(·,4)。
- result1/2：工作表 温度/水分浓度；result3/4：Sheet1。A1=`时间\\到药材中心的距离`。
- result4：列 0..1.9 + 「药材表面」；域外单元格留空（None→空）。
- 数据行 ≤1,048,575（含表头 ≤1,048,576），超限抛错，不截断、不改模板。
- verify_workbook：读回校验表名/A1/表头/A 列连续无重复/行数/掩码。
"""
from __future__ import annotations

from pathlib import Path

import numpy as np
import openpyxl
from openpyxl.cell import WriteOnlyCell

from .postprocess import round4

EXCEL_MAX_ROWS = 1048576
NUMFMT = "0.0000"


def _num_cell(ws, value):
    """数值型单元格（round4 + 0.0000 格式）；None → 空单元格。"""
    if value is None:
        return None
    c = WriteOnlyCell(ws, value=round4(value))
    c.number_format = NUMFMT
    return c


def _check_row_budget(n_data_rows: int, path):
    if n_data_rows + 1 > EXCEL_MAX_ROWS:
        raise ValueError(f"{path}: 数据行 {n_data_rows} 超出上限 {EXCEL_MAX_ROWS - 1}")


def write_result12(path, t_array, C_grid, T_grid, cols_cm, a1_text):
    """result1/2：两工作表（温度、水分浓度）。C_grid/T_grid 形状 (len(t), len(cols))。"""
    path = Path(path)
    t_array = np.asarray(t_array)
    _check_row_budget(len(t_array), path)
    wb = openpyxl.Workbook(write_only=True)
    for sheet_name, grid, is_temp in (("温度", T_grid, True), ("水分浓度", C_grid, False)):
        ws = wb.create_sheet(sheet_name)
        ws.append([a1_text] + [float(c) for c in cols_cm])
        for i, t in enumerate(t_array):
            row = [int(round(t))] + [_num_cell(ws, grid[i][j]) for j in range(len(cols_cm))]
            ws.append(row)
    wb.save(path)


def write_result34(path, t_array, value_grid, cols_cm, a1_text, *,
                   sheet_name="Sheet1", surface_header=None, surface_values=None):
    """result3/4：单工作表（水分浓度）。value_grid 可含 None（域外留空）。

    result4 传 surface_header='药材表面' 与 surface_values（长度 len(t)）。
    """
    path = Path(path)
    t_array = np.asarray(t_array)
    _check_row_budget(len(t_array), path)
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet(sheet_name)
    header = [a1_text] + [float(c) for c in cols_cm]
    if surface_header is not None:
        header.append(surface_header)
    ws.append(header)
    for i, t in enumerate(t_array):
        row = [int(round(t))] + [_num_cell(ws, value_grid[i][j]) for j in range(len(cols_cm))]
        if surface_header is not None:
            row.append(_num_cell(ws, surface_values[i]))
        ws.append(row)
    wb.save(path)


def verify_workbook(path, *, expected_sheets, a1_text, expected_cols, n_data_rows,
                    t_start, t_step, surface_header=None, mask=None):
    """写后读回校验（V-9）。mask[i][j]=True 表域内（应非空）。返回检查结果 dict。"""
    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True)
    issues = []
    if list(wb.sheetnames) != list(expected_sheets):
        issues.append(f"工作表名 {wb.sheetnames} != {expected_sheets}")
    for sn in wb.sheetnames:
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        header = rows[0]
        if header[0] != a1_text:
            issues.append(f"{sn}: A1={header[0]!r} != {a1_text!r}")
        exp_header = tuple([a1_text] + list(expected_cols) +
                           ([surface_header] if surface_header else []))
        got_header = tuple(header)
        if len(got_header) != len(exp_header):
            issues.append(f"{sn}: 表头列数 {len(got_header)} != {len(exp_header)}")
        # A 列连续整数、无重复
        a_col = [r[0] for r in rows[1:]]
        if len(a_col) != n_data_rows:
            issues.append(f"{sn}: 数据行 {len(a_col)} != {n_data_rows}")
        expected_a = list(range(int(t_start), int(t_start) + n_data_rows * int(t_step), int(t_step)))
        if a_col != expected_a:
            issues.append(f"{sn}: A 列非预期连续整数（首 {a_col[:3]} 末 {a_col[-3:]}）")
        if len(set(a_col)) != len(a_col):
            issues.append(f"{sn}: A 列有重复")
        # 掩码一致性（result4）
        if mask is not None:
            for i, r in enumerate(rows[1:]):
                # 数据列（去掉 A 列）
                data = r[1:1 + len(expected_cols)]
                for j, v in enumerate(data):
                    is_none = v is None
                    should_be_inside = bool(mask[i][j])
                    if should_be_inside == is_none:
                        issues.append(f"{sn}: 行{i} 列{j} 掩码不一致 (inside={should_be_inside}, empty={is_none})")
                        break
    wb.close()
    return {"ok": len(issues) == 0, "issues": issues}
