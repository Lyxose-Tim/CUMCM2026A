"""data_io.py —— 原始附件只读加载 + 环境/半径驱动函数（SI 单位）。

对齐《A题_建模方案.md》§6.0.5、D03、D07、B12：
- 附件 1（时间 s / 温度 °C / 水分浓度 kg/kg）：0–14400 s 原始节点分段线性插值；
  t>14400 s 取 [10800,14400] s **61 个原始样本的算术均值**（运行时全精度计算）；
  14400 s 取原始节点、右侧显式切换为常值 → 登记求解断点。
- `smooth121` 情景：内点 y_i←(y_{i-1}+2y_i+y_{i+1})/4（i=1..239），两端原值，再线性插值；
  **外推均值不变**（仍用原始窗口）。
- 灵敏度：外推段 T_air ±0.39 °C、C_env ±0.0011（**仅作用于外推段**）；`hold_last` 用末点。
- 附件 2（时间 s / 半径 cm）：分段线性插值；t>72 h 保持 1.198 cm 并标注外推。

原始附件一律只读，不删改。温度对外以 K 返回（供 D(C,T)）。
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import numpy as np
import openpyxl

KELVIN = 273.15


# --------------------------------------------------------------------------
# 原始加载（只读）
# --------------------------------------------------------------------------
def load_attachment1(path: str | Path) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """返回 (t_s, T_degC, C_kgkg)；附件 1，Sheet1，跳过表头。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    t, T, C = [], [], []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        if row[0] is None:
            continue
        t.append(float(row[0]))
        T.append(float(row[1]))
        C.append(float(row[2]))
    wb.close()
    return np.array(t), np.array(T), np.array(C)


def load_attachment2(path: str | Path) -> tuple[np.ndarray, np.ndarray]:
    """返回 (t_s, R_cm)；附件 2，Sheet1，跳过表头。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    t, R = [], []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        if row[0] is None:
            continue
        t.append(float(row[0]))
        R.append(float(row[1]))
    wb.close()
    return np.array(t), np.array(R)


def _smooth121(y: np.ndarray) -> np.ndarray:
    """内点 [1,2,1]/4 平滑，两端保持原值（§6.0.5）。"""
    out = y.copy()
    out[1:-1] = 0.25 * (y[:-2] + 2.0 * y[1:-1] + y[2:])
    return out


# --------------------------------------------------------------------------
# 环境驱动函数
# --------------------------------------------------------------------------
@dataclass
class EnvFunctions:
    """环境时序驱动。t 单位 s；T_air 返回 K，C_env 无量纲。"""
    t_nodes: np.ndarray            # 用于插值的时间节点（0–14400 s）
    T_nodes_K: np.ndarray          # 对应温度 (K)（可能已平滑）
    C_nodes: np.ndarray            # 对应含水率（可能已平滑）
    t_switch: float                # 断点（14400 s）
    T_const_K: float               # 外推常值温度 (K)
    C_const: float                 # 外推常值含水率
    breakpoints: tuple[float, ...] = (14400.0,)

    def T_air_K(self, t):
        t = np.asarray(t, dtype=np.float64)
        interp = np.interp(t, self.t_nodes, self.T_nodes_K)
        out = np.where(t <= self.t_switch, interp, self.T_const_K)
        return out if out.shape else float(out)

    def C_env(self, t):
        t = np.asarray(t, dtype=np.float64)
        interp = np.interp(t, self.t_nodes, self.C_nodes)
        out = np.where(t <= self.t_switch, interp, self.C_const)
        return out if out.shape else float(out)


SCENARIO_ENV_PARAMS: dict[str, dict] = {
    "base": {},
    "smooth121": {"interp": "smooth121"},
    "hold_last": {"extrapolation": "hold_last"},
    "dT_air+": {"dT_extrap": +0.39},
    "dT_air-": {"dT_extrap": -0.39},
    "dC_env+": {"dC_extrap": +0.0011},
    "dC_env-": {"dC_extrap": -0.0011},
}


def make_env_functions(cfg, scenario: str = "base") -> EnvFunctions:
    """构造环境驱动。scenario 见 SCENARIO_ENV_PARAMS。"""
    if scenario not in SCENARIO_ENV_PARAMS:
        raise KeyError(f"未知环境情景 {scenario!r}")
    params = SCENARIO_ENV_PARAMS[scenario]
    interp = params.get("interp", cfg.raw["air"].get("interp", "linear"))
    extrapolation = params.get("extrapolation", cfg.raw["air"].get("extrapolation", "hold_window_mean"))
    dT_extrap = params.get("dT_extrap", 0.0)
    dC_extrap = params.get("dC_extrap", 0.0)

    t, T_C, C = load_attachment1(cfg.air_file())
    t_switch = float(cfg.breakpoints_s[0])          # 14400 s

    # 外推窗口 [10800,14400] 的 61 点均值（原始数据，全精度；平滑不改此均值）
    w0, w1 = cfg.air_window_s
    win = (t >= w0) & (t <= w1)
    T_win_mean_C = float(np.mean(T_C[win]))
    C_win_mean = float(np.mean(C[win]))

    # 插值节点（可选平滑，仅改原数据段）
    T_nodes_C = T_C.copy()
    C_nodes = C.copy()
    if interp == "smooth121":
        T_nodes_C = _smooth121(T_C)
        C_nodes = _smooth121(C)

    # 外推常值
    if extrapolation == "hold_last":
        # 末点保持 (50.165, 0.04986)
        last = t == t_switch
        T_const_C = float(T_C[last][-1])
        C_const = float(C[last][-1])
    else:  # hold_window_mean
        T_const_C = T_win_mean_C
        C_const = C_win_mean
    T_const_C += dT_extrap
    C_const += dC_extrap

    return EnvFunctions(
        t_nodes=t,
        T_nodes_K=T_nodes_C + KELVIN,
        C_nodes=C_nodes,
        t_switch=t_switch,
        T_const_K=T_const_C + KELVIN,
        C_const=C_const,
        breakpoints=(t_switch,),
    )


# --------------------------------------------------------------------------
# 半径驱动函数（Q4）
# --------------------------------------------------------------------------
@dataclass
class RadiusFunction:
    """R(t)：附件 2 分段线性插值 + t>72 h 保持末值。t 单位 s，返回 m。"""
    t_nodes: np.ndarray
    R_nodes_m: np.ndarray
    t_last: float
    R_last_m: float

    def R(self, t):
        t = np.asarray(t, dtype=np.float64)
        out = np.interp(t, self.t_nodes, self.R_nodes_m)  # 越界自动取端点值 → 即 hold_last
        return out if out.shape else float(out)

    def is_extrapolated(self, t) -> bool:
        return float(np.max(np.asarray(t))) > self.t_last


def make_radius_function(cfg) -> RadiusFunction:
    t, R_cm = load_attachment2(cfg.radius_file())
    R_m = R_cm / 100.0
    return RadiusFunction(
        t_nodes=t,
        R_nodes_m=R_m,
        t_last=float(t[-1]),
        R_last_m=float(R_m[-1]),
    )
